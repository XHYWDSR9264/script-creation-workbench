from __future__ import annotations

import hashlib
import base64
import json
import re
import sqlite3
import uuid
from datetime import datetime, timezone
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import unquote, urlparse, quote
from urllib.request import Request, urlopen
from urllib.error import HTTPError, URLError

ROOT = Path(__file__).resolve().parent
DB_PATH = ROOT / "workbench.sqlite3"
STORAGE_ROOT = ROOT / "data" / "projects"
ASSET_TYPES = ("synopsis", "characters", "world", "beat_matrix")
GATES = (("G0", "立项与参数"), ("G1", "方案与卡点"), ("G2", "第1—3集校准"), ("G3", "批次创作"), ("G4", "内容终审"), ("G5", "朱雀同版检测"), ("G6", "交付与归档"))


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def db() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode=WAL")
    conn.execute("PRAGMA foreign_keys=ON")
    return conn


def seed_workspace(conn: sqlite3.Connection, project_id: str, stamp: str, profile: dict | None = None) -> None:
    p = profile or {}
    conn.execute("""INSERT OR IGNORE INTO project_profiles
        (project_id,region,medium,genre,total_episodes,submission_start,submission_end,opening_template,current_stage,progress,status,updated_at,duration_min_seconds,duration_max_seconds,calibration_threshold,quality_threshold,zhuque_threshold,one_scene_each,zhuque_required)
        VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
        (project_id, p.get("region", "国内"), p.get("medium", "AI漫剧"), p.get("genre", "待设定"),
         int(p.get("totalEpisodes", 60)), int(p.get("submissionStart", 1)), int(p.get("submissionEnd", 10)),
         p.get("openingTemplate", "老王模板"), "G0 立项与参数", 10, "待立项", stamp,
         int(p.get("durationMin", 120)), int(p.get("durationMax", 180)), int(p.get("calibrationThreshold", 85)),
         int(p.get("qualityThreshold", 90)), float(p.get("zhuqueThreshold", 85)), 0 if p.get("oneSceneEach") is False else 1,
         0 if p.get("zhuqueRequired") is False else 1))
    for index, (code, title) in enumerate(GATES):
        conn.execute("INSERT OR IGNORE INTO gates(project_id,gate_code,title,status,note,updated_at) VALUES(?,?,?,?,?,?)",
                     (project_id, code, title, "current" if index == 0 else "pending", "等待参数确认" if index == 0 else "未开始", stamp))
    for asset_type in ASSET_TYPES:
        conn.execute("INSERT OR IGNORE INTO story_assets(project_id,asset_type,content,version,updated_at) VALUES(?,?,?,1,?)", (project_id, asset_type, "", stamp))


def init_db() -> None:
    with db() as conn:
        conn.executescript((ROOT / "schema.sql").read_text(encoding="utf-8"))
        def ensure_column(table: str, column: str, definition: str) -> None:
            names = {row[1] for row in conn.execute(f"PRAGMA table_info({table})")}
            if column not in names: conn.execute(f"ALTER TABLE {table} ADD COLUMN {column} {definition}")
        for column, definition in (
            ("duration_min_seconds", "INTEGER NOT NULL DEFAULT 120"), ("duration_max_seconds", "INTEGER NOT NULL DEFAULT 180"),
            ("calibration_threshold", "INTEGER NOT NULL DEFAULT 85"), ("quality_threshold", "INTEGER NOT NULL DEFAULT 90"),
            ("zhuque_threshold", "REAL NOT NULL DEFAULT 85"), ("one_scene_each", "INTEGER NOT NULL DEFAULT 1"),
            ("zhuque_required", "INTEGER NOT NULL DEFAULT 1"), ("origin_research_id", "TEXT"),
            ("origin_candidate_id", "TEXT")):
            ensure_column("project_profiles", column, definition)
        for column, definition in (
            ("audit_type", "TEXT NOT NULL DEFAULT 'structural'"), ("dimensions_json", "TEXT NOT NULL DEFAULT '{}'"),
            ("reverse_checks_json", "TEXT NOT NULL DEFAULT '{}'"), ("hard_errors_json", "TEXT NOT NULL DEFAULT '[]'"),
            ("core_floor_pass", "INTEGER NOT NULL DEFAULT 0")):
            ensure_column("audit_reports", column, definition)
        ensure_column("task_runs", "retry_of", "TEXT")
        for row in conn.execute("SELECT id,updated_at FROM projects").fetchall():
            seed_workspace(conn, row["id"], row["updated_at"])


def project_summary(row: sqlite3.Row) -> dict:
    return {"id": row["id"], "name": row["name"], "route": row["route"], "note": row["note"], "updatedAt": row["updated_at"]}


def get_detail(conn: sqlite3.Connection, project_id: str) -> dict | None:
    project = conn.execute("SELECT * FROM projects WHERE id=?", (project_id,)).fetchone()
    if not project:
        return None
    profile = conn.execute("SELECT * FROM project_profiles WHERE project_id=?", (project_id,)).fetchone()
    assets = conn.execute("SELECT asset_type,content,version,updated_at FROM story_assets WHERE project_id=?", (project_id,)).fetchall()
    versions = conn.execute("SELECT id,label,range_start,range_end,sha256,source,status,created_at FROM draft_versions WHERE project_id=? ORDER BY created_at DESC,rowid DESC", (project_id,)).fetchall()
    gates = conn.execute("SELECT gate_code,title,status,note,updated_at FROM gates WHERE project_id=? ORDER BY gate_code", (project_id,)).fetchall()
    events = conn.execute("SELECT kind,message,created_at FROM activity_logs WHERE project_id=? ORDER BY id DESC LIMIT 50", (project_id,)).fetchall()
    sources = conn.execute("SELECT id,name,mime,byte_size,sha256,extraction_status,created_at FROM source_documents WHERE project_id=? ORDER BY created_at DESC", (project_id,)).fetchall()
    audits = conn.execute("SELECT id,version_id,score,status,summary,findings_json,sha256,created_at,audit_type,dimensions_json,reverse_checks_json,hard_errors_json,core_floor_pass FROM audit_reports WHERE project_id=? ORDER BY created_at DESC LIMIT 20", (project_id,)).fetchall()
    detectors = conn.execute("SELECT id,version_id,human_score,suspected_score,ai_score,report_name,sha256,status,created_at FROM detector_reports WHERE project_id=? ORDER BY created_at DESC LIMIT 20", (project_id,)).fetchall()
    runs = conn.execute("SELECT id,task_type,instruction,status,stage,progress,output_preview,error,version_id,retry_of,created_at,updated_at FROM task_runs WHERE project_id=? ORDER BY created_at DESC LIMIT 30", (project_id,)).fetchall()
    approvals = conn.execute("SELECT id,gate_code,decision,note,version_id,sha256,actor,created_at FROM approval_records WHERE project_id=? ORDER BY created_at DESC LIMIT 30", (project_id,)).fetchall()
    continuity = conn.execute("SELECT id,category,subject,state,first_episode,last_episode,source_version_id,status,notes,created_at,updated_at FROM continuity_entries WHERE project_id=? ORDER BY category,subject", (project_id,)).fetchall()
    quality_reviews = conn.execute("SELECT id,version_id,score,note,actor,created_at FROM quality_reviews WHERE project_id=? ORDER BY created_at DESC LIMIT 30", (project_id,)).fetchall()
    latest_version_id = versions[0]["id"] if versions else None
    result = {"project": project_summary(project), "profile": dict(profile) if profile else None,
            "assets": {r["asset_type"]: {"content": r["content"], "version": r["version"], "updatedAt": r["updated_at"]} for r in assets},
            "versions": [dict(r) for r in versions], "gates": [dict(r) for r in gates], "events": [dict(r) for r in events], "sources": [dict(r) for r in sources],
            "audits": [{**dict(r), "is_current": r["version_id"] == latest_version_id} for r in audits],
            "detectors": [{**dict(r), "is_current": r["version_id"] == latest_version_id} for r in detectors],
            "runs": [dict(r) for r in runs], "approvals": [dict(r) for r in approvals], "qualityReviews": [dict(r) for r in quality_reviews],
            "continuity": [{**dict(r), "is_current": not r["source_version_id"] or r["source_version_id"] == latest_version_id} for r in continuity],
            "latestVersionId": latest_version_id}
    result["workflow"] = build_workflow(result)
    return result


def has_real_value(value: object) -> bool:
    text = str(value or "").strip()
    return bool(text and text not in {"待设定", "未设定", "无"})


def covers(version: dict | sqlite3.Row | None, start: int, end: int) -> bool:
    return bool(version and int(version["range_start"]) <= start and int(version["range_end"]) >= end)


def asset_ready(asset_type: str, value: object) -> bool:
    text = str(value or "").strip(); minimum = {"synopsis": 100, "characters": 120, "world": 80, "beat_matrix": 150}.get(asset_type, 80)
    return len(text) >= minimum and not re.search(r"待填写|待补充|TODO|TBD", text, re.I)


def build_workflow(data: dict) -> dict:
    profile = data.get("profile") or {}; gates = data.get("gates") or []; assets = data.get("assets") or {}; versions = data.get("versions") or []; latest = versions[0] if versions else None; latest_id = data.get("latestVersionId")
    start, end = int(profile.get("submission_start", 1)), int(profile.get("submission_end", 10)); calibration_end = min(end, start + 2)
    missing = [x for x in ("synopsis", "characters", "world", "beat_matrix") if not asset_ready(x, (assets.get(x) or {}).get("content", ""))]
    audits, detectors = data.get("audits") or [], data.get("detectors") or []; passed = {x["gate_code"] for x in gates if x["status"] == "passed"}
    calibration_version = next((x for x in versions if covers(x, start, calibration_end)), None)
    structural = next((x for x in audits if x.get("is_current") and x["version_id"] == (calibration_version or {}).get("id") and x.get("audit_type") == "structural" and x.get("status") == "通过"), None)
    deep = any(x.get("is_current") and x["version_id"] == latest_id and x.get("audit_type") == "deep_ai" and x.get("status") == "通过" and int(x.get("core_floor_pass", 0)) == 1 for x in audits)
    quality_threshold = float(profile.get("quality_threshold", 90)); quality = next((x for x in data.get("qualityReviews", []) if x["version_id"] == latest_id and float(x["score"]) >= quality_threshold), None)
    detector_required = int(profile.get("zhuque_required", 1)) != 0; zhuque_threshold = float(profile.get("zhuque_threshold", 85)); detector = any(x.get("is_current") and x["version_id"] == latest_id and x.get("status") == "通过" and float(x["human_score"]) >= zhuque_threshold for x in detectors)
    checks = [
        {"code": "G0", "title": "立项与参数", "ready": all((has_real_value(data.get("project", {}).get("name")), has_real_value(data.get("project", {}).get("route")), has_real_value(profile.get("region")), has_real_value(profile.get("medium")), has_real_value(profile.get("genre")), int(profile.get("total_episodes", 0)) > 0, has_real_value(profile.get("opening_template")))), "reason": "立项参数已齐" if has_real_value(profile.get("genre")) else "补齐地区、形式、题材、总集数和开头模板等立项参数"},
        {"code": "G1", "title": "方案与卡点", "ready": not missing, "reason": f"项目资料待补：{'、'.join(missing)}" if missing else "四项项目资料已齐"},
        {"code": "G2", "title": "第1—3集校准", "ready": bool(calibration_version and structural), "reason": ("校准稿结构审计已通过" if calibration_version and structural else ("先对校准稿运行结构审计" if calibration_version else "先保存覆盖校准范围的正文版本"))},
        {"code": "G3", "title": "批次创作", "ready": bool(latest and covers(latest, start, end) and "G2" in passed), "reason": ("先批准G2校准闸门" if "G2" not in passed else ("正文版本尚未覆盖本次提交范围" if not latest or not covers(latest, start, end) else "提交范围正文已齐"))},
        {"code": "G4", "title": "内容终审", "ready": bool(deep and quality and "G3" in passed), "reason": ("先批准G3批次创作" if "G3" not in passed else ("先对当前正文运行通过的AI内部深审" if not deep else (f"先为当前正文录入不低于{quality_threshold:g}分的人工内容质量评分" if not quality else "深审与内容质量评分已齐")))},
        {"code": "G5", "title": "朱雀同版检测", "ready": bool((not detector_required or detector) and "G4" in passed), "reason": ("先批准G4内容终审" if "G4" not in passed else ("项目已关闭朱雀必检" if not detector_required else ("当前正文朱雀报告已达标" if detector else f"先保存当前正文实际朱雀报告，人工特征需达到{zhuque_threshold:g}%")))},
        {"code": "G6", "title": "交付与归档", "ready": bool("G4" in passed and deep and quality and (not detector_required or ("G5" in passed and detector)) and latest and covers(latest, start, end)), "reason": ("先完成当前正文的G4内容终审" if "G4" not in passed or not deep or not quality else ("先完成当前正文的G5朱雀同版检测" if detector_required and ("G5" not in passed or not detector) else ("可提交最终交付" if latest and covers(latest, start, end) else "正文尚未覆盖本次提交范围")))},
    ]
    current = next((x for x in gates if x["status"] == "current"), None) or next((x for x in gates if x["status"] != "passed"), None) or (gates[-1] if gates else None); current_check = next((x for x in checks if x["code"] == (current or {}).get("gate_code")), checks[0]); passed_count = sum(1 for x in gates if x["status"] == "passed"); complete = passed_count == len(GATES)
    return {"currentGate": (current or {}).get("gate_code", "G0"), "currentStage": "G6 已归档" if complete else f"{(current or {}).get('gate_code', 'G0')} {(current or {}).get('title', '立项与参数')}", "progress": 100 if complete else round(passed_count / len(GATES) * 100), "checks": checks, "complete": complete, "nextAction": {"gate": current_check["code"], "title": current_check["title"], "ready": complete or current_check["ready"], "message": ("G0—G6已全部通过，当前项目已归档" if complete else (f"可以处理{current_check['code']}：{current_check['title']}" if current_check["ready"] else f"{current_check['code']}暂不能批准：{current_check['reason']}"))}, "missingAssets": missing, "latestVersionCoversSubmission": bool(latest and covers(latest, start, end)), "calibrationEnd": calibration_end}


def invalidate_review_gates(conn: sqlite3.Connection, project_id: str, stamp: str) -> None:
    g3 = conn.execute("SELECT status FROM gates WHERE project_id=? AND gate_code='G3'", (project_id,)).fetchone()
    conn.execute("UPDATE gates SET status='pending',note='正文已更新，旧终审证据失效',updated_at=? WHERE project_id=? AND gate_code IN ('G4','G5','G6')", (stamp, project_id))
    if g3 and g3["status"] == "passed":
        conn.execute("UPDATE gates SET status='current',note='等待当前正文重新终审',updated_at=? WHERE project_id=? AND gate_code='G4'", (stamp, project_id))
        conn.execute("UPDATE project_profiles SET current_stage='G4 内容终审',progress=50,status='等待重新终审',updated_at=? WHERE project_id=?", (stamp, project_id))


def parse_episode_number(token: str) -> int:
    if token.isdigit(): return int(token)
    nums = {"零": 0, "〇": 0, "一": 1, "二": 2, "两": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9}
    if token == "十": return 10
    if "百" in token:
        left, right = (token.split("百", 1) + [""])[:2]
        return nums.get(left, 1) * 100 + parse_episode_number(right or "0")
    if "十" in token:
        left, right = (token.split("十", 1) + [""])[:2]
        return (nums.get(left, 1) if left else 1) * 10 + (nums.get(right, 0) if right else 0)
    value = 0
    for char in token: value = value * 10 + nums.get(char, 0)
    return value


def audit_text(content: str, start: int, end: int) -> dict:
    findings: list[dict] = []
    def add(level: str, code: str, message: str) -> None: findings.append({"level": level, "code": code, "message": message})
    episodes = [parse_episode_number(x) for x in re.findall(r"^\s*第\s*([0-9一二三四五六七八九十百〇零两]+)\s*集\s*$", content, re.M)]
    missing = [x for x in range(start, end + 1) if x not in episodes]
    outside = sorted(set(x for x in episodes if x < start or x > end))
    if missing: add("error", "missing_episode", "缺少集标题：第" + "、".join(map(str, missing)) + "集")
    if outside: add("error", "range_leak", "正文混入提交范围外集数：第" + "、".join(map(str, outside)) + "集")
    if len(set(episodes)) != len(episodes): add("error", "duplicate_episode", "存在重复集标题")
    scenes = len(re.findall(r"^\s*\d+\s*[-－—]\s*\d+\s+", content, re.M))
    if not scenes: add("error", "scene_header", "未识别到规范场次标题")
    if re.search(r"&#x?[0-9a-f]+;", content, re.I): add("error", "html_entity", "正文含HTML实体乱码")
    if "\\n" in content: add("error", "literal_newline", "正文含字面量\\n")
    if re.search(r"【\s*黑幕\s*】", content): add("warning", "black_screen", "正文含【黑幕】字样，请确认交付格式是否允许")
    if re.search(r"（\s*(本次提交|检测范围)[^）]*）", content): add("warning", "delivery_marker", "正文含检测或提交范围说明")
    for left, right in (("（", "）"), ("【", "】"), ("《", "》")):
        if content.count(left) != content.count(right): add("error", "bracket_balance", f"{left}{right}数量不平衡：{content.count(left)}/{content.count(right)}")
    counts: dict[str, int] = {}
    for line in (x.strip() for x in content.splitlines()):
        if len(line) >= 8: counts[line] = counts.get(line, 0) + 1
    duplicates = [x for x, count in counts.items() if count >= 3]
    if duplicates: add("warning", "repeated_lines", f"发现{len(duplicates)}条重复三次以上的长句")
    if len(content.strip()) < 800: add("warning", "short_content", "正文过短，无法完成有效结构审计")
    errors = sum(1 for x in findings if x["level"] == "error"); warnings = len(findings) - errors; score = max(0, 100 - errors * 12 - warnings * 4)
    return {"score": score, "status": "通过" if score >= 90 and errors == 0 else "需修复", "summary": f"识别{len(episodes)}个集标题、{scenes}个场次；{errors}项错误、{warnings}项提醒。", "findings": findings}


DIMENSIONS = {"premise": 15, "causality": 15, "characters": 20, "dialogue": 15,
              "pacing": 10, "emotion": 10, "performance": 10, "continuity": 5}


def parse_json_object(text: str) -> dict:
    clean = re.sub(r"^```(?:json)?\s*", "", str(text or "").strip(), flags=re.I)
    clean = re.sub(r"```\s*$", "", clean).strip()
    start, end = clean.find("{"), clean.rfind("}")
    if start < 0 or end <= start:
        raise ValueError("模型未返回可解析的JSON对象")
    return json.loads(clean[start:end + 1])


def research_prompt(title: str, brief: str, reference_titles: str, market_notes: str, count: int) -> str:
    return f"""你是短剧选题研发主编。请根据用户提供的创作方向、近期剧名信号和可核验备注，提出{count}个彼此显著不同、可进入正式立项的原创短剧方案。

调研主题：{title}
用户思路：{brief or '未提供'}
近期剧名/参考标题（只代表标题与题材信号，不代表你知道其剧情、热度或收益）：
{reference_titles or '未提供'}
可核验市场数据或用户备注：
{market_notes or '未提供'}

硬性规则：
1. 不得从剧名臆造原作剧情、签约状态、播放量、收益或排名；没有数据就明确按“标题信号”分析。
2. 推荐必须是新项目，不复刻参考标题，不只换人名、职业或时代；人物关系、核心机制、冲突场域和结局至少三项不同。
3. 各方案题材、职业、情绪引擎、视觉奇观要拉开差异；剧名不要全部使用“两段式反转标题”。
4. 建议体量限定30—60集，并结合事件容量判断，不机械统一为60集。
5. score是内部立项推荐分，不是平台通过率；按题眼15、冲突20、人物15、持续性20、差异化15、可视化15合计100分保守评分。
6. 只输出一个JSON对象，不要Markdown，不要解释性前后缀。

JSON结构：{{"analysis_summary":"说明本轮仅用了哪些信号、哪些事实不能判断","candidates":[{{"title":"","genre":"","hook":"","core_conflict":"","innovation":"","episode_recommendation":50,"score":90,"risks":""}}]}}"""


def normalize_research(raw: dict, count: int, reference_titles: str) -> dict:
    refs = set()
    for value in re.split(r"\r?\n|[；;]", reference_titles or ""):
        value = re.sub(r"^\s*\d+[.、）)]?\s*", "", value).replace("《", "").replace("》", "").strip()
        if value: refs.add(value)
    candidates, seen = [], set()
    for source in raw.get("candidates") or []:
        title = str(source.get("title") or "").replace("《", "").replace("》", "").strip(); key = title.lower()
        if not title or len(title) > 80 or key in seen or title in refs: continue
        genre = str(source.get("genre") or "").strip(); hook = str(source.get("hook") or "").strip(); core = str(source.get("core_conflict") or "").strip(); innovation = str(source.get("innovation") or "").strip(); risks = str(source.get("risks") or "").strip()
        if not genre or len(hook) < 18 or len(core) < 18 or len(innovation) < 12: continue
        try: episodes = round(float(source.get("episode_recommendation") or 50))
        except (TypeError, ValueError): episodes = 50
        try: score = float(source.get("score") or 0)
        except (TypeError, ValueError): score = 0
        seen.add(key); candidates.append({"title": title, "genre": genre, "hook": hook, "core_conflict": core, "innovation": innovation, "episode_recommendation": max(30, min(60, episodes)), "score": max(0, min(100, score)), "risks": risks or "需在G1阶段继续核查同质化、专业常识与长线容量"})
        if len(candidates) >= count: break
    if len(candidates) < min(3, count): raise ValueError(f"有效推荐不足3个（仅解析到{len(candidates)}个），请重试或补充更明确的思路")
    return {"analysisSummary": str(raw.get("analysis_summary") or "本轮只依据用户提供的创作方向与标题信号形成推荐；未提供的数据不作事实判断。").strip(), "candidates": candidates}


def normalize_deep_audit(raw: dict, threshold: float) -> dict:
    dimensions, total = {}, 0.0
    for key, maximum in DIMENSIONS.items():
        source = (raw.get("dimensions") or {}).get(key) or {}
        try: score = float(source.get("score", 0))
        except (TypeError, ValueError): score = 0
        score = max(0, min(maximum, score)); total += score
        dimensions[key] = {"score": score, "max": maximum,
                           "evidence": list(source.get("evidence") or [])[:6],
                           "deduction": str(source.get("deduction") or "")}
    total = round(total, 1)
    hard_errors = list(raw.get("hard_errors") or [])[:50]
    reverse_source = raw.get("reverse_checks") if isinstance(raw.get("reverse_checks"), dict) else {}
    reverse_checks = {}
    for key in ("prop_chain", "permission_license", "adjacent_transition", "presence_speaker", "packaging_identity"):
        source = reverse_source.get(key) or {}; state = source.get("status") if source.get("status") in {"pass", "fail", "not_applicable"} else "fail"
        reverse_checks[key] = {"status": state, "evidence": list(source.get("evidence") or [])[:8],
                               "problem": str(source.get("problem") or ("模型漏交该反查表" if key not in reverse_source else ""))}
    core_floor = dimensions["causality"]["score"] >= 12 and dimensions["characters"]["score"] >= 16 and dimensions["dialogue"]["score"] >= 12
    reverse_pass = all(x["status"] in {"pass", "not_applicable"} for x in reverse_checks.values())
    status = "通过" if total >= threshold and core_floor and reverse_pass and not hard_errors else "需修复"
    findings = [{"level": "error", "code": str(x.get("type") or "hard_error"),
                 "message": f"{x.get('episode','')}{('/' + str(x.get('scene'))) if x.get('scene') else ''} {x.get('evidence') or x.get('problem') or '硬错误'}".strip()}
                for x in hard_errors if isinstance(x, dict)]
    findings.extend({"level": "error", "code": f"reverse_{key}", "message": value["problem"] or f"{key}反查未通过"}
                    for key, value in reverse_checks.items() if value["status"] == "fail")
    return {"score": total, "status": status,
            "summary": str(raw.get("summary") or f"AI内部八维深审{total}分；硬错误{len(hard_errors)}项。"),
            "findings": findings, "dimensions": dimensions, "reverseChecks": reverse_checks,
            "hardErrors": hard_errors, "coreFloorPass": core_floor}


def deep_audit_prompt(data: dict, version: sqlite3.Row) -> str:
    return f'''你是严格的中文短剧SOP内部审读员。只审读给定版本，不改稿，不因目标分倒填。必须输出单个JSON对象，不要Markdown。
项目：{data['project']['name']}
路线：{data['project']['route']}
总集数：{data.get('profile', {}).get('total_episodes', 60)}
本次正文：第{version['range_start']}—{version['range_end']}集
门槛：{data.get('profile', {}).get('quality_threshold', 90)}

按八维评分：premise满15；causality满15且底线12；characters满20且底线16；dialogue满15且底线12；pacing满10；emotion满10；performance满10；continuity满5。每维返回score、evidence数组（必须含集号/场次/短原文锚点）、deduction。
按五张反查表返回reverse_checks：prop_chain、permission_license、adjacent_transition、presence_speaker、packaging_identity；每项返回status(pass/fail/not_applicable)、evidence数组、problem。另查人物所知、时间空间、伤势、金额期限、重生/异能边界、反派利益、专业常识。
硬错误放hard_errors数组，每项含episode、scene、type、evidence、minimal_fix、affected_later。summary必须说明最大优点和首要缺陷。
JSON结构：{{"dimensions":{{"premise":{{"score":0,"evidence":[],"deduction":""}}}},"reverse_checks":{{}},"hard_errors":[],"summary":""}}

项目资料：
梗概：{data.get('assets', {}).get('synopsis', {}).get('content', '无')}
人物：{data.get('assets', {}).get('characters', {}).get('content', '无')}
世界观：{data.get('assets', {}).get('world', {}).get('content', '无')}
卡点：{data.get('assets', {}).get('beat_matrix', {}).get('content', '无')}

正文：
{version['content']}'''


def analyze_script(content: str) -> dict:
    episodes, characters, locations = [], {}, {}
    current_episode = current_scene = None
    for index, raw_line in enumerate(content.splitlines(), 1):
        line = raw_line.strip()
        ep = re.match(r"^第\s*([0-9一二三四五六七八九十百〇零两]+)\s*集$", line)
        if ep:
            current_episode = {"number": parse_episode_number(ep.group(1)), "line": index, "scenes": []}
            episodes.append(current_episode); current_scene = None; continue
        scene = re.match(r"^(\d+)\s*[-－—]\s*(\d+)\s+(.+)$", line)
        if scene:
            current_scene = {"episode": current_episode["number"] if current_episode else int(scene.group(1)),
                             "number": int(scene.group(2)), "heading": line, "line": index, "characters": []}
            if current_episode is None:
                current_episode = {"number": int(scene.group(1)), "line": index, "scenes": []}; episodes.append(current_episode)
            current_episode["scenes"].append(current_scene)
            bits = scene.group(3).split(); location = " ".join(bits[2:]) if len(bits) > 2 else scene.group(3)
            locations[location] = locations.get(location, 0) + 1; continue
        cast = re.match(r"^人物[：:]\s*(.+)$", line)
        if cast and current_scene:
            names = [x.strip() for x in re.split(r"[、，,/／]", cast.group(1)) if x.strip()]
            current_scene["characters"] = names
            for name in names: characters[name] = characters.get(name, 0) + 1
    return {"episodes": episodes, "episodeCount": len(episodes),
            "sceneCount": sum(len(x["scenes"]) for x in episodes),
            "characters": [{"name": k, "scenes": v} for k, v in sorted(characters.items(), key=lambda x: -x[1])],
            "locations": [{"name": k, "scenes": v} for k, v in sorted(locations.items(), key=lambda x: -x[1])]}


def model_endpoint(base: str, wire: str) -> str:
    root = base.rstrip("/")
    if not re.match(r"^https?://", root, re.I): raise ValueError("Base URL必须以http://或https://开头")
    suffix = "/chat/completions" if wire == "chat" else "/responses"
    return root if root.endswith(suffix) else root + suffix


def call_model(config: dict, prompt: str) -> str:
    base_url = str(config.get("baseUrl", "")).strip(); api_key = str(config.get("apiKey", "")).strip(); model = str(config.get("model", "")).strip(); wire = "chat" if config.get("wireApi") == "chat" else "responses"
    if not base_url or not api_key or not model: raise ValueError("请先填写Base URL、API Key和模型名")
    payload = {"model": model, "messages": [{"role": "user", "content": prompt}], "temperature": 0.7} if wire == "chat" else {"model": model, "input": prompt, "reasoning": {"effort": str(config.get("reasoning", "high"))}}
    req = Request(model_endpoint(base_url, wire), data=json.dumps(payload, ensure_ascii=False).encode(), method="POST", headers={"Content-Type": "application/json", "Authorization": "Bearer " + api_key})
    try:
        raw = urlopen(req, timeout=180).read().decode("utf-8", errors="replace")
    except HTTPError as exc:
        msg = exc.read().decode("utf-8", errors="replace")[:300]; raise ValueError(f"模型接口HTTP {exc.code}：{msg}") from exc
    except URLError as exc: raise ValueError(f"模型接口连接失败：{exc.reason}") from exc
    data = json.loads(raw)
    if wire == "chat": output = (((data.get("choices") or [{}])[0].get("message") or {}).get("content"))
    else:
        output = data.get("output_text")
        if not output:
            output = "\n".join(c.get("text", "") for item in data.get("output", []) for c in item.get("content", []) if c.get("type") == "output_text")
    if not output: raise ValueError("接口返回成功，但没有读取到文本内容")
    return str(output)


def model_prompt(data: dict, task_type: str, instruction: str) -> str:
    a = data.get("assets", {}); p = data.get("profile") or {}; sources = "、".join(f"{x['name']}（{x['extraction_status']}）" for x in data.get("sources", [])) or "无"
    source_text = "\n\n".join(f"【{x['name']}】\n{x['extracted_text']}" for x in data.get("sourceTexts", []) if x.get("extracted_text"))[:120000] or "无可用文本素材"
    return f"""你是短剧生产工作台的内容引擎。当前任务类型：{task_type}。
项目：{data['project']['name']}
路线：{data['project']['route']}
参数：{p.get('region','')} / {p.get('medium','')} / 总{p.get('total_episodes',60)}集 / 提交{p.get('submission_start',1)}-{p.get('submission_end',10)}集 / {p.get('opening_template','')}
题材：{p.get('genre','')}
素材清单：{sources}
可用素材正文：
{source_text}
现有梗概：{a.get('synopsis',{}).get('content','无')}
现有人物：{a.get('characters',{}).get('content','无')}
现有世界观：{a.get('world',{}).get('content','无')}
现有卡点：{a.get('beat_matrix',{}).get('content','无')}

用户指令：{instruction}

必须尊重提交范围，不得把范围外分集写进剧情梗概；生成完成不等于人工审校通过。"""


def export_text(data: dict, version: sqlite3.Row | None) -> str:
    a = data.get("assets", {}); p = data["project"]; profile = data.get("profile") or {}; value = lambda key: a.get(key, {}).get("content", "")
    return f"《{p['name']}》\n\n生产路线：{p['route']}\n地区：{profile.get('region','')}\n形式：{profile.get('medium','')}\n总集数：{profile.get('total_episodes','')}\n开头模板：{profile.get('opening_template','')}\n\n一、剧情梗概\n{value('synopsis')}\n\n二、人物小传\n{value('characters')}\n\n三、世界观设定\n{value('world')}\n\n四、卡点矩阵\n{value('beat_matrix')}\n\n五、剧本正文\n{version['content'] if version else '尚未选择正文版本'}\n\n正文SHA-256：{version['sha256'] if version else '无'}\n"


def extract_source_text(path: Path) -> tuple[str, str]:
    ext = path.suffix.lower()
    if ext in {".txt", ".md", ".markdown", ".json", ".csv", ".tsv"}:
        return path.read_text(encoding="utf-8", errors="replace"), "已提取"
    if ext == ".docx":
        try:
            from docx import Document
            doc = Document(path)
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
            for table in doc.tables:
                text += "\n" + "\n".join("\t".join(cell.text for cell in row.cells) for row in table.rows)
            return text.strip(), "已提取"
        except Exception:
            return "", "待解析"
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
            return "\n".join((p.extract_text() or "") for p in PdfReader(path).pages).strip(), "已提取"
        except Exception:
            return "", "待解析"
    return "", "待解析"


class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(ROOT), **kwargs)

    def log_message(self, fmt, *args):
        return

    def send_json(self, payload: dict, status: int = 200) -> None:
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status); self.send_header("Content-Type", "application/json; charset=utf-8"); self.send_header("Cache-Control", "no-store"); self.send_header("Content-Length", str(len(body))); self.end_headers(); self.wfile.write(body)

    def send_download(self, body: bytes, content_type: str, filename: str) -> None:
        self.send_response(200); self.send_header("Content-Type", content_type); self.send_header("Content-Disposition", "attachment; filename*=UTF-8''" + quote(filename)); self.send_header("Content-Length", str(len(body))); self.end_headers(); self.wfile.write(body)

    def read_json(self) -> dict:
        length = int(self.headers.get("Content-Length", "0"))
        if length > 15_000_000: raise ValueError("请求内容过大")
        return json.loads(self.rfile.read(length).decode("utf-8") if length else "{}")

    @staticmethod
    def parts(path: str) -> list[str]:
        return [unquote(x) for x in path.split("/") if x]

    def do_GET(self):
        path = urlparse(self.path).path.rstrip("/") or "/"; parts = self.parts(path)
        if path == "/api/health": self.send_json({"ok": True, "service": "剧本创作总控台", "storage": "sqlite", "version": "0.8.0"}); return
        if path == "/api/research":
            with db() as conn: rows = conn.execute("SELECT r.*,COUNT(c.id) AS candidate_count FROM research_sessions r LEFT JOIN research_candidates c ON c.session_id=r.id GROUP BY r.id ORDER BY r.updated_at DESC LIMIT 50").fetchall()
            self.send_json({"sessions": [dict(r) for r in rows]}); return
        if len(parts) == 3 and parts[:2] == ["api", "research"]:
            with db() as conn:
                session = conn.execute("SELECT * FROM research_sessions WHERE id=?", (parts[2],)).fetchone()
                candidates = conn.execute("SELECT id,session_id,title,genre,hook,core_conflict,innovation,episode_recommendation,score,risks,created_at FROM research_candidates WHERE session_id=? ORDER BY score DESC,created_at", (parts[2],)).fetchall() if session else []
            self.send_json({"session": dict(session), "candidates": [dict(x) for x in candidates]} if session else {"error": "调研记录不存在"}, 200 if session else 404); return
        if path == "/api/projects":
            with db() as conn: rows = conn.execute("SELECT * FROM projects ORDER BY updated_at DESC").fetchall()
            self.send_json({"projects": [project_summary(r) for r in rows]}); return
        if len(parts) == 3 and parts[:2] == ["api", "projects"]:
            with db() as conn: detail = get_detail(conn, parts[2])
            self.send_json({"detail": detail} if detail else {"error": "项目不存在"}, 200 if detail else 404); return
        if len(parts) == 4 and parts[:2] == ["api", "versions"] and parts[3] == "content":
            with db() as conn: row = conn.execute("SELECT * FROM draft_versions WHERE id=?", (parts[2],)).fetchone()
            self.send_json({"version": dict(row)} if row else {"error": "版本不存在"}, 200 if row else 404); return
        if len(parts) == 4 and parts[:2] == ["api", "versions"] and parts[3] == "analysis":
            with db() as conn: row = conn.execute("SELECT id,label,content,sha256 FROM draft_versions WHERE id=?", (parts[2],)).fetchone()
            self.send_json({"version": {"id": row["id"], "label": row["label"], "sha256": row["sha256"]}, "analysis": analyze_script(row["content"])} if row else {"error": "版本不存在"}, 200 if row else 404); return
        if len(parts) == 4 and parts[:2] == ["api", "sources"] and parts[3] == "content":
            with db() as conn: row = conn.execute("SELECT id,name,mime,byte_size,sha256,extracted_text,extraction_status,created_at FROM source_documents WHERE id=?", (parts[2],)).fetchone()
            self.send_json({"source": dict(row)} if row else {"error": "素材不存在"}, 200 if row else 404); return
        if len(parts) == 4 and parts[:2] == ["api", "projects"] and parts[3] == "export":
            query = urlparse(self.path).query; params = dict(x.split("=", 1) if "=" in x else (x, "") for x in query.split("&") if x); version_id = unquote(params.get("versionId", "")); fmt = params.get("format", "txt")
            with db() as conn:
                detail = get_detail(conn, parts[2]); version = conn.execute("SELECT * FROM draft_versions WHERE id=? AND project_id=?", (version_id, parts[2])).fetchone() if version_id else conn.execute("SELECT * FROM draft_versions WHERE project_id=? ORDER BY created_at DESC,rowid DESC LIMIT 1", (parts[2],)).fetchone()
            if not detail: self.send_json({"error": "项目不存在"}, 404); return
            safe = re.sub(r'[\\/:*?"<>|]', "_", detail["project"]["name"])
            if fmt == "json": self.send_download(json.dumps({"detail": detail, "selectedVersion": dict(version) if version else None}, ensure_ascii=False, indent=2).encode(), "application/json; charset=utf-8", safe + ".json"); return
            self.send_download(export_text(detail, version).encode(), "text/plain; charset=utf-8", safe + ".txt"); return
        super().do_GET()

    def do_POST(self):
        path = urlparse(self.path).path.rstrip("/"); parts = self.parts(path)
        try:
            if path == "/api/research/generate":
                data = self.read_json(); title = str(data.get("title", "")).strip(); brief = str(data.get("brief", "")).strip(); reference_titles = str(data.get("referenceTitles", "")).strip(); market_notes = str(data.get("marketNotes", "")).strip()
                try: count = int(data.get("count", 5))
                except (TypeError, ValueError): count = 0
                if not title or len(title) > 120: self.send_json({"error": "调研主题不能为空且不超过120字"}, 400); return
                if count < 3 or count > 10: self.send_json({"error": "推荐数量必须是3—10之间的整数"}, 400); return
                if not brief and not reference_titles and not market_notes: self.send_json({"error": "请至少提供创作思路、近期剧名或市场备注中的一项"}, 400); return
                session_id = "rs-" + uuid.uuid4().hex[:12]; stamp = now_iso()
                with db() as conn: conn.execute("INSERT INTO research_sessions(id,title,brief,reference_titles,market_notes,requested_count,status,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?)", (session_id, title, brief, reference_titles, market_notes, count, "running", stamp, stamp))
                try: output = call_model(data.get("config") or {}, research_prompt(title, brief, reference_titles, market_notes, count))
                except ValueError as exc:
                    with db() as conn: conn.execute("UPDATE research_sessions SET status=?,error=?,updated_at=? WHERE id=?", ("failed", str(exc), now_iso(), session_id))
                    self.send_json({"error": str(exc), "sessionId": session_id}, 502); return
                try: normalized = normalize_research(parse_json_object(output), count, reference_titles)
                except (ValueError, json.JSONDecodeError) as exc:
                    with db() as conn: conn.execute("UPDATE research_sessions SET status=?,error=?,updated_at=? WHERE id=?", ("failed", str(exc), now_iso(), session_id))
                    self.send_json({"error": str(exc), "sessionId": session_id}, 422); return
                done = now_iso(); candidate_rows = []
                with db() as conn:
                    conn.execute("UPDATE research_sessions SET analysis_summary=?,status=?,error='',updated_at=? WHERE id=?", (normalized["analysisSummary"], "ready", done, session_id))
                    for candidate in normalized["candidates"]:
                        candidate_id = "rc-" + uuid.uuid4().hex[:12]
                        conn.execute("INSERT INTO research_candidates(id,session_id,title,genre,hook,core_conflict,innovation,episode_recommendation,score,risks,raw_json,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)", (candidate_id, session_id, candidate["title"], candidate["genre"], candidate["hook"], candidate["core_conflict"], candidate["innovation"], candidate["episode_recommendation"], candidate["score"], candidate["risks"], json.dumps(candidate, ensure_ascii=False), done))
                    candidate_rows = conn.execute("SELECT id,session_id,title,genre,hook,core_conflict,innovation,episode_recommendation,score,risks,created_at FROM research_candidates WHERE session_id=? ORDER BY score DESC,created_at", (session_id,)).fetchall()
                self.send_json({"session": {"id": session_id, "title": title, "brief": brief, "reference_titles": reference_titles, "market_notes": market_notes, "requested_count": count, "analysis_summary": normalized["analysisSummary"], "status": "ready", "created_at": stamp, "updated_at": done}, "candidates": [dict(x) for x in candidate_rows]}, 201); return
            if len(parts) == 4 and parts[:2] == ["api", "research"] and parts[3] == "project":
                data = self.read_json()
                with db() as conn: candidate = conn.execute("SELECT c.*,r.id AS research_id,r.title AS research_title FROM research_candidates c JOIN research_sessions r ON r.id=c.session_id WHERE c.id=?", (parts[2],)).fetchone()
                if not candidate: self.send_json({"error": "推荐方案不存在"}, 404); return
                name = str(data.get("name") or candidate["title"]).strip()
                try: total = int(data.get("totalEpisodes") or candidate["episode_recommendation"] or 50); submission_end = int(data.get("submissionEnd") or 10); min_duration = int(data.get("durationMin") or 120); max_duration = int(data.get("durationMax") or 180); quality = float(data.get("qualityThreshold") or 90); zhuque = float(data.get("zhuqueThreshold") or 85)
                except (TypeError, ValueError): total, submission_end, min_duration, max_duration, quality, zhuque = 0, 0, 0, 0, -1, -1
                if not name or len(name) > 120: self.send_json({"error": "项目名称不能为空且不超过120字"}, 400); return
                if total < 10 or total > 200 or submission_end < 1 or submission_end > total: self.send_json({"error": "总集数或提交范围无效"}, 400); return
                if min_duration < 30 or max_duration < min_duration or max_duration > 600 or not (0 <= quality <= 100 and 0 <= zhuque <= 100): self.send_json({"error": "立项门槛或单集时长无效"}, 400); return
                project_data = {**data, "totalEpisodes": total, "submissionStart": 1, "submissionEnd": submission_end, "durationMin": min_duration, "durationMax": max_duration, "qualityThreshold": quality, "zhuqueThreshold": zhuque, "route": data.get("route") or "完全原创 / 市场参考", "genre": data.get("genre") or candidate["genre"], "openingTemplate": data.get("openingTemplate") or "老王模板"}
                project_id = "p-" + uuid.uuid4().hex[:12]; stamp = now_iso(); seed = f"选题种子（需在G1扩写并人工确认）\n一句话钩子：{candidate['hook']}\n核心冲突：{candidate['core_conflict']}\n创新点：{candidate['innovation']}\n风险提示：{candidate['risks']}"
                with db() as conn:
                    conn.execute("INSERT INTO projects(id,name,route,note,created_at,updated_at) VALUES(?,?,?,?,?,?)", (project_id, name, project_data["route"], "调研推荐立项 · 待G0确认", stamp, stamp)); seed_workspace(conn, project_id, stamp, project_data)
                    conn.execute("UPDATE project_profiles SET origin_research_id=?,origin_candidate_id=? WHERE project_id=?", (candidate["research_id"], candidate["id"], project_id)); conn.execute("UPDATE story_assets SET content=?,version=version+1,updated_at=? WHERE project_id=? AND asset_type='synopsis'", (seed, stamp, project_id)); conn.execute("INSERT INTO activity_logs(project_id,kind,message,created_at) VALUES(?,?,?,?)", (project_id, "research", f"从调研「{candidate['research_title']}」推荐方案建立项目；仍需完成G0参数确认与G1完整方案", stamp))
                self.send_json({"project": {"id": project_id, "name": name, "route": project_data["route"], "note": "调研推荐立项 · 待G0确认", "updatedAt": stamp}}, 201); return
            if path == "/api/model/test":
                data = self.read_json()
                try: output = call_model(data, "仅回复“连接成功”，不要补充其他内容。")
                except ValueError as exc: self.send_json({"error": str(exc)}, 502); return
                self.send_json({"ok": True, "output": output}); return
            if path == "/api/projects":
                data = self.read_json(); name = str(data.get("name", "")).strip(); route = str(data.get("route", "完全原创 / 市场参考")).strip()
                if not name or len(name) > 120: self.send_json({"error": "项目名称不能为空且不超过120字"}, 400); return
                min_duration, max_duration = int(data.get("durationMin", 120)), int(data.get("durationMax", 180)); quality, zhuque = float(data.get("qualityThreshold", 90)), float(data.get("zhuqueThreshold", 85)); total = int(data.get("totalEpisodes", 60)); submission_start = int(data.get("submissionStart", 1)); submission_end = int(data.get("submissionEnd", 10))
                if total < 1 or submission_start < 1 or submission_end < submission_start or submission_end > total: self.send_json({"error": "总集数和本次提交范围必须有效，且提交范围不能超过总集数"}, 400); return
                if min_duration < 30 or max_duration < min_duration or max_duration > 600: self.send_json({"error": "单集时长范围无效"}, 400); return
                if not (0 <= quality <= 100 and 0 <= zhuque <= 100): self.send_json({"error": "质量和朱雀门槛必须在0—100之间"}, 400); return
                project_id = "p-" + uuid.uuid4().hex[:12]; stamp = now_iso(); note = "新项目 · 待立项"
                with db() as conn:
                    conn.execute("INSERT INTO projects(id,name,route,note,created_at,updated_at) VALUES(?,?,?,?,?,?)", (project_id, name, route, note, stamp, stamp)); seed_workspace(conn, project_id, stamp, data)
                    conn.execute("INSERT INTO activity_logs(project_id,kind,message,created_at) VALUES(?,?,?,?)", (project_id, "project", "创建项目：" + name, stamp)); row = conn.execute("SELECT * FROM projects WHERE id=?", (project_id,)).fetchone()
                self.send_json({"project": project_summary(row)}, 201); return
            if len(parts) == 4 and parts[:2] == ["api", "projects"] and parts[3] == "events":
                data = self.read_json(); message = str(data.get("message", "")).strip(); kind = str(data.get("kind", "task"))[:30]
                if not message: self.send_json({"error": "事件内容不能为空"}, 400); return
                stamp = now_iso()
                with db() as conn:
                    if not conn.execute("SELECT 1 FROM projects WHERE id=?", (parts[2],)).fetchone(): self.send_json({"error": "项目不存在"}, 404); return
                    conn.execute("INSERT INTO activity_logs(project_id,kind,message,created_at) VALUES(?,?,?,?)", (parts[2], kind, message, stamp)); conn.execute("UPDATE projects SET updated_at=? WHERE id=?", (stamp, parts[2]))
                self.send_json({"ok": True, "createdAt": stamp}); return
            if len(parts) == 4 and parts[:2] == ["api", "projects"] and parts[3] == "versions":
                data = self.read_json(); content = str(data.get("content", "")); label = str(data.get("label", "")).strip(); start, end = int(data.get("rangeStart", 1)), int(data.get("rangeEnd", 3))
                if not label or not content.strip() or start < 1 or end < start: self.send_json({"error": "版本名称、正文和集数范围必须有效"}, 400); return
                version_id = "v-" + uuid.uuid4().hex[:12]; stamp = now_iso(); digest = hashlib.sha256(content.encode("utf-8")).hexdigest(); source = str(data.get("source", "人工编辑"))
                with db() as conn:
                    profile = conn.execute("SELECT total_episodes FROM project_profiles WHERE project_id=?", (parts[2],)).fetchone()
                    if not profile: self.send_json({"error": "项目不存在"}, 404); return
                    if end > int(profile["total_episodes"] or 60): self.send_json({"error": f"正文范围不能超过项目总集数（{profile['total_episodes']}集）"}, 400); return
                    conn.execute("INSERT INTO draft_versions(id,project_id,label,range_start,range_end,content,sha256,source,status,created_at) VALUES(?,?,?,?,?,?,?,?,?,?)", (version_id, parts[2], label, start, end, content, digest, source, "已保存", stamp))
                    invalidate_review_gates(conn, parts[2], stamp)
                    conn.execute("INSERT INTO activity_logs(project_id,kind,message,created_at) VALUES(?,?,?,?)", (parts[2], "version", f"保存正文版本 {label}（第{start}—{end}集）", stamp)); conn.execute("UPDATE projects SET updated_at=? WHERE id=?", (stamp, parts[2]))
                self.send_json({"version": {"id": version_id, "label": label, "range_start": start, "range_end": end, "sha256": digest, "source": source, "status": "已保存", "created_at": stamp}}, 201); return
            if len(parts) == 4 and parts[:2] == ["api", "projects"] and parts[3] == "audits":
                data = self.read_json(); project_id = parts[2]; version_id = str(data.get("versionId", ""))
                with db() as conn: version = conn.execute("SELECT * FROM draft_versions WHERE id=? AND project_id=?", (version_id, project_id)).fetchone()
                if not version: self.send_json({"error": "请选择当前项目的正文版本"}, 400); return
                report = audit_text(version["content"], version["range_start"], version["range_end"]); report_id = "a-" + uuid.uuid4().hex[:12]; stamp = now_iso()
                with db() as conn:
                    conn.execute("INSERT INTO audit_reports(id,project_id,version_id,score,status,summary,findings_json,sha256,created_at,audit_type,dimensions_json,reverse_checks_json,hard_errors_json,core_floor_pass) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)", (report_id, project_id, version_id, report["score"], report["status"], report["summary"], json.dumps(report["findings"], ensure_ascii=False), version["sha256"], stamp, "structural", "{}", "{}", "[]", 1 if report["status"] == "通过" else 0))
                    conn.execute("INSERT INTO activity_logs(project_id,kind,message,created_at) VALUES(?,?,?,?)", (project_id, "audit", f"结构审计：{report['score']}分，{report['status']}", stamp))
                self.send_json({"report": {"id": report_id, "version_id": version_id, "sha256": version["sha256"], "created_at": stamp, **report}}, 201); return
            if len(parts) == 5 and parts[:2] == ["api", "projects"] and parts[3:] == ["audits", "deep"]:
                data = self.read_json(); project_id = parts[2]; version_id = str(data.get("versionId", ""))
                with db() as conn:
                    version = conn.execute("SELECT * FROM draft_versions WHERE id=? AND project_id=?", (version_id, project_id)).fetchone(); detail = get_detail(conn, project_id)
                if not version or not detail: self.send_json({"error": "请选择当前项目的正文版本"}, 400); return
                run_id = "r-" + uuid.uuid4().hex[:12]; stamp = now_iso()
                with db() as conn: conn.execute("INSERT INTO task_runs(id,project_id,task_type,instruction,status,stage,progress,version_id,retry_of,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)", (run_id, project_id, "deep_audit", f"深审版本 {version['label']}", "running", "八维评分与五表反查", 35, version_id, str(data.get("retryOf") or "") or None, stamp, stamp))
                try: output = call_model(data.get("config") or {}, deep_audit_prompt(detail, version))
                except ValueError as exc:
                    with db() as conn: conn.execute("UPDATE task_runs SET status=?,stage=?,progress=?,error=?,updated_at=? WHERE id=?", ("failed", "模型调用失败", 100, str(exc), now_iso(), run_id))
                    self.send_json({"error": str(exc), "runId": run_id}, 502); return
                try: report = normalize_deep_audit(parse_json_object(output), float((detail.get("profile") or {}).get("quality_threshold", 90)))
                except (ValueError, json.JSONDecodeError) as exc:
                    with db() as conn: conn.execute("UPDATE task_runs SET status=?,stage=?,progress=?,error=?,output_preview=?,updated_at=? WHERE id=?", ("failed", "结果解析失败", 100, str(exc), output[:3500], now_iso(), run_id))
                    self.send_json({"error": str(exc), "runId": run_id}, 422); return
                report_id = "a-" + uuid.uuid4().hex[:12]; done = now_iso()
                with db() as conn:
                    conn.execute("INSERT INTO audit_reports(id,project_id,version_id,score,status,summary,findings_json,sha256,created_at,audit_type,dimensions_json,reverse_checks_json,hard_errors_json,core_floor_pass) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)", (report_id, project_id, version_id, report["score"], report["status"], report["summary"], json.dumps(report["findings"], ensure_ascii=False), version["sha256"], done, "deep_ai", json.dumps(report["dimensions"], ensure_ascii=False), json.dumps(report["reverseChecks"], ensure_ascii=False), json.dumps(report["hardErrors"], ensure_ascii=False), 1 if report["coreFloorPass"] else 0))
                    conn.execute("UPDATE task_runs SET status=?,stage=?,progress=?,output_preview=?,updated_at=? WHERE id=?", ("succeeded", "等待人工确认", 100, report["summary"][:3500], done, run_id))
                    conn.execute("INSERT INTO activity_logs(project_id,kind,message,created_at) VALUES(?,?,?,?)", (project_id, "audit", f"AI内部深审：{report['score']}分，{report['status']}；硬错误{len(report['hardErrors'])}项", done))
                self.send_json({"report": {"id": report_id, "version_id": version_id, "sha256": version["sha256"], "created_at": done, "audit_type": "deep_ai", **report}, "runId": run_id}, 201); return
            if len(parts) == 4 and parts[:2] == ["api", "projects"] and parts[3] == "quality":
                data = self.read_json(); project_id = parts[2]; version_id = str(data.get("versionId", "")); note = str(data.get("note", "")).strip()
                try: score = float(data.get("score"))
                except (TypeError, ValueError): self.send_json({"error": "内容质量评分必须在0—100之间"}, 400); return
                if score < 0 or score > 100: self.send_json({"error": "内容质量评分必须在0—100之间"}, 400); return
                with db() as conn:
                    project = conn.execute("SELECT id FROM projects WHERE id=?", (project_id,)).fetchone()
                    version = conn.execute("SELECT id,sha256 FROM draft_versions WHERE id=? AND project_id=?", (version_id, project_id)).fetchone()
                    latest = conn.execute("SELECT id FROM draft_versions WHERE project_id=? ORDER BY created_at DESC,rowid DESC LIMIT 1", (project_id,)).fetchone()
                if not project: self.send_json({"error": "项目不存在"}, 404); return
                if not version: self.send_json({"error": "请选择当前项目的正文版本"}, 400); return
                if not latest or latest["id"] != version_id: self.send_json({"error": "内容质量评分只能绑定当前最新正文"}, 409); return
                review_id = "q-" + uuid.uuid4().hex[:12]; stamp = now_iso()
                with db() as conn:
                    conn.execute("INSERT INTO quality_reviews(id,project_id,version_id,score,note,actor,created_at) VALUES(?,?,?,?,?,?,?)", (review_id, project_id, version_id, score, note, "operator", stamp))
                    conn.execute("INSERT INTO activity_logs(project_id,kind,message,created_at) VALUES(?,?,?,?)", (project_id, "quality", f"人工内容质量评分：{score:g}分", stamp))
                self.send_json({"review": {"id": review_id, "project_id": project_id, "version_id": version_id, "score": score, "note": note, "actor": "operator", "created_at": stamp, "sha256": version["sha256"]}}, 201); return
            if len(parts) == 4 and parts[:2] == ["api", "projects"] and parts[3] == "detectors":
                data = self.read_json(); project_id = parts[2]; version_id = str(data.get("versionId", "")); human = float(data.get("humanScore", -1)); suspected = float(data.get("suspectedScore", 0)); ai = float(data.get("aiScore", 0)); name = str(data.get("reportName", "")).strip(); encoded = str(data.get("dataBase64", ""))
                if any(x < 0 or x > 100 for x in (human, suspected, ai)): self.send_json({"error": "检测指标必须在0—100之间"}, 400); return
                if not name or not encoded: self.send_json({"error": "必须上传朱雀报告截图"}, 400); return
                try: raw = base64.b64decode(encoded, validate=True)
                except Exception: self.send_json({"error": "报告截图编码无效"}, 400); return
                if len(raw) > 5 * 1024 * 1024: self.send_json({"error": "报告截图不能超过5MB"}, 413); return
                if not (raw.startswith(b"\x89PNG\r\n\x1a\n") or raw.startswith(b"\xff\xd8\xff")): self.send_json({"error": "报告文件必须是真实PNG或JPEG图片"}, 400); return
                with db() as conn:
                    version = conn.execute("SELECT * FROM draft_versions WHERE id=? AND project_id=?", (version_id, project_id)).fetchone()
                    profile = conn.execute("SELECT zhuque_threshold FROM project_profiles WHERE project_id=?", (project_id,)).fetchone()
                if not version: self.send_json({"error": "请选择当前项目的正文版本"}, 400); return
                threshold = float(profile["zhuque_threshold"] if profile else 85)
                report_id = "z-" + uuid.uuid4().hex[:12]; stamp = now_iso(); safe_name = "".join(c if c.isalnum() or c in "._-" else "_" for c in name); folder = STORAGE_ROOT / project_id / "detectors" / report_id; folder.mkdir(parents=True, exist_ok=True); path = folder / safe_name; path.write_bytes(raw); object_key = str(path.relative_to(ROOT)).replace("\\", "/"); status = "通过" if human >= threshold else "未通过"
                with db() as conn:
                    conn.execute("INSERT INTO detector_reports(id,project_id,version_id,human_score,suspected_score,ai_score,report_object_key,report_name,sha256,status,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)", (report_id, project_id, version_id, human, suspected, ai, object_key, name, version["sha256"], status, stamp))
                    conn.execute("INSERT INTO activity_logs(project_id,kind,message,created_at) VALUES(?,?,?,?)", (project_id, "detector", f"朱雀报告：人工特征{human}%，{status}", stamp))
                self.send_json({"report": {"id": report_id, "version_id": version_id, "human_score": human, "suspected_score": suspected, "ai_score": ai, "report_name": name, "sha256": version["sha256"], "status": status, "created_at": stamp}}, 201); return
            if len(parts) == 5 and parts[:2] == ["api", "projects"] and parts[3:] == ["model", "run"]:
                data = self.read_json(); project_id = parts[2]; task_type = str(data.get("taskType", "reply")); instruction = str(data.get("instruction", "")).strip()
                if not instruction: self.send_json({"error": "任务内容不能为空"}, 400); return
                if task_type not in set(ASSET_TYPES) | {"draft", "reply"}: self.send_json({"error": "模型任务类型无效"}, 400); return
                with db() as conn:
                    detail = get_detail(conn, project_id)
                    if detail: detail["sourceTexts"] = [dict(x) for x in conn.execute("SELECT name,extracted_text FROM source_documents WHERE project_id=? AND extraction_status=? ORDER BY created_at", (project_id, "已提取")).fetchall()]
                if not detail: self.send_json({"error": "项目不存在"}, 404); return
                run_id = "r-" + uuid.uuid4().hex[:12]; started = now_iso()
                with db() as conn: conn.execute("INSERT INTO task_runs(id,project_id,task_type,instruction,status,stage,progress,retry_of,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?)", (run_id, project_id, task_type, instruction, "running", "模型生成", 35, str(data.get("retryOf") or "") or None, started, started))
                try: output = call_model(data.get("config") or {}, model_prompt(detail, task_type, instruction))
                except ValueError as exc:
                    with db() as conn: conn.execute("UPDATE task_runs SET status=?,stage=?,progress=?,error=?,updated_at=? WHERE id=?", ("failed", "模型调用失败", 100, str(exc), now_iso(), run_id))
                    self.send_json({"error": str(exc), "runId": run_id}, 502); return
                stamp = now_iso()
                version_id = None
                with db() as conn:
                    if task_type in ASSET_TYPES:
                        conn.execute("""INSERT INTO story_assets(project_id,asset_type,content,version,updated_at) VALUES(?,?,?,1,?) ON CONFLICT(project_id,asset_type) DO UPDATE SET content=excluded.content,version=story_assets.version+1,updated_at=excluded.updated_at""", (project_id, task_type, output, stamp))
                    elif task_type == "draft":
                        version_id = "v-" + uuid.uuid4().hex[:12]; digest = hashlib.sha256(output.encode()).hexdigest(); start = int(data.get("rangeStart", detail.get("profile", {}).get("submission_start", 1))); end = int(data.get("rangeEnd", detail.get("profile", {}).get("submission_end", 10)))
                        total = int((detail.get("profile") or {}).get("total_episodes", 60))
                        if start < 1 or end < start or end > total:
                            message = f"模型正文范围无效，不能超过项目总集数（{total}集）"
                            conn.execute("UPDATE task_runs SET status=?,stage=?,progress=?,error=?,output_preview=?,updated_at=? WHERE id=?", ("failed", "范围校验失败", 100, message, output[:3500], stamp, run_id))
                            self.send_json({"error": message, "runId": run_id}, 400); return
                        conn.execute("INSERT INTO draft_versions(id,project_id,label,range_start,range_end,content,sha256,source,status,created_at) VALUES(?,?,?,?,?,?,?,?,?,?)", (version_id, project_id, str(data.get("label", f"模型稿 {stamp[:16]}")), start, end, output, digest, "模型生成", "待人工审校", stamp))
                        invalidate_review_gates(conn, project_id, stamp)
                    conn.execute("UPDATE task_runs SET status=?,stage=?,progress=?,output_preview=?,version_id=?,updated_at=? WHERE id=?", ("succeeded", "等待人工审校" if task_type == "draft" else "已保存", 100, output[:3500], version_id, stamp, run_id))
                    message = f"模型任务完成：{task_type}" + (("\n" + output[:3500]) if task_type == "reply" else "")
                    conn.execute("INSERT INTO activity_logs(project_id,kind,message,created_at) VALUES(?,?,?,?)", (project_id, "model", message, stamp)); conn.execute("UPDATE projects SET updated_at=? WHERE id=?", (stamp, project_id))
                self.send_json({"ok": True, "output": output, "taskType": task_type, "runId": run_id, "versionId": version_id}); return
            if len(parts) == 5 and parts[:2] == ["api", "projects"] and parts[3] == "gates":
                data = self.read_json(); project_id, gate_code = parts[2], parts[4]; decision = str(data.get("decision", "")); note = str(data.get("note", "")).strip(); version_id = str(data.get("versionId", ""))
                if gate_code not in {x[0] for x in GATES} or decision not in {"approve", "reject"}: self.send_json({"error": "闸门或决定无效"}, 400); return
                with db() as conn:
                    detail = get_detail(conn, project_id); version = conn.execute("SELECT id,sha256 FROM draft_versions WHERE id=? AND project_id=?", (version_id, project_id)).fetchone() if version_id else None
                if not detail: self.send_json({"error": "项目不存在"}, 404); return
                gate_index = next(i for i, x in enumerate(GATES) if x[0] == gate_code)
                current_gate = next((x for x in detail["gates"] if x["status"] == "current"), None)
                if current_gate and current_gate["gate_code"] != gate_code: self.send_json({"error": f"当前只能处理{current_gate['gate_code']} {current_gate['title']}"}, 409); return
                if decision == "approve" and next((x for x in detail["gates"] if x["gate_code"] == gate_code and x["status"] == "passed"), None): self.send_json({"error": "该闸门已经批准，不能重复提交"}, 409); return
                if decision == "approve" and gate_index > 0:
                    prior = next((x for x in detail["gates"][:gate_index] if x["status"] != "passed"), None)
                    if prior: self.send_json({"error": f"必须先通过{prior['gate_code']} {prior['title']}"}, 409); return
                if decision == "approve" and gate_code in {"G2", "G3", "G4", "G5", "G6"} and not version: self.send_json({"error": "该阶段批准必须绑定正文版本"}, 400); return
                if decision == "approve" and version and version["id"] != detail.get("latestVersionId"): self.send_json({"error": "只能批准当前最新正文版本"}, 409); return
                if decision == "approve" and gate_code in {"G0", "G1", "G2", "G3"}:
                    check = next((x for x in detail["workflow"]["checks"] if x["code"] == gate_code), None)
                    if not check or not check["ready"]: self.send_json({"error": (check or {}).get("reason", "当前阶段前置条件尚未满足")}, 409); return
                if decision == "approve" and gate_code == "G4":
                    ok = any(x["audit_type"] == "deep_ai" and x["version_id"] == version_id and x["status"] == "通过" and int(x["core_floor_pass"]) == 1 for x in detail["audits"])
                    if not ok: self.send_json({"error": "G4需要当前版本AI内部深审达到项目门槛、核心底线通过且硬错误为0"}, 409); return
                    quality_threshold = float((detail.get("profile") or {}).get("quality_threshold", 90))
                    quality_ok = any(x["version_id"] == version_id and float(x["score"]) >= quality_threshold for x in detail.get("qualityReviews", []))
                    if not quality_ok: self.send_json({"error": f"G4还需要当前正文人工内容质量评分达到{quality_threshold:g}分"}, 409); return
                if decision == "approve" and gate_code == "G5" and int((detail.get("profile") or {}).get("zhuque_required", 1)) != 0:
                    threshold = float((detail.get("profile") or {}).get("zhuque_threshold", 85)); ok = any(x["version_id"] == version_id and x["status"] == "通过" and float(x["human_score"]) >= threshold for x in detail["detectors"])
                    if not ok: self.send_json({"error": f"G5需要当前版本朱雀人工特征达到{threshold:g}%并保存真实报告"}, 409); return
                if decision == "approve" and gate_code == "G6":
                    check = next((x for x in detail["workflow"]["checks"] if x["code"] == "G6"), None)
                    if not check or not check["ready"]: self.send_json({"error": (check or {}).get("reason", "G6需要当前正文通过内容终审和适用的朱雀门禁")}, 409); return
                approval_id = "ap-" + uuid.uuid4().hex[:12]; stamp = now_iso(); status = "passed" if decision == "approve" else "blocked"; sha = version["sha256"] if version else ""
                with db() as conn:
                    conn.execute("INSERT INTO approval_records(id,project_id,gate_code,decision,note,version_id,sha256,actor,created_at) VALUES(?,?,?,?,?,?,?,?,?)", (approval_id, project_id, gate_code, decision, note, version_id or None, sha, "operator", stamp))
                    conn.execute("UPDATE gates SET status=?,note=?,updated_at=? WHERE project_id=? AND gate_code=?", (status, note or status, stamp, project_id, gate_code))
                    if decision == "approve":
                        idx = gate_index
                        if idx < len(GATES) - 1: conn.execute("UPDATE gates SET status=?,note=?,updated_at=? WHERE project_id=? AND gate_code=? AND status<>?", ("current", "等待处理", stamp, project_id, GATES[idx + 1][0], "passed"))
                        next_gate = GATES[min(idx + 1, len(GATES) - 1)]
                        conn.execute("UPDATE project_profiles SET current_stage=?,progress=?,status=?,updated_at=? WHERE project_id=?", (f"{next_gate[0]} {next_gate[1]}", 100 if gate_code == "G6" else round(((idx + 1) / len(GATES)) * 100), "已归档" if gate_code == "G6" else "等待人工处理", stamp, project_id))
                    else:
                        idx = gate_index
                        conn.execute("UPDATE project_profiles SET current_stage=?,status=?,updated_at=? WHERE project_id=?", (f"{gate_code} {GATES[idx][1]}", "已阻塞", stamp, project_id))
                    conn.execute("INSERT INTO activity_logs(project_id,kind,message,created_at) VALUES(?,?,?,?)", (project_id, "approval", f"{gate_code}{'批准' if decision == 'approve' else '驳回'}：{note or '无备注'}", stamp))
                self.send_json({"approval": {"id": approval_id, "gate_code": gate_code, "decision": decision, "note": note, "version_id": version_id or None, "sha256": sha, "created_at": stamp}}, 201); return
            if len(parts) == 4 and parts[:2] == ["api", "projects"] and parts[3] == "continuity":
                data = self.read_json(); project_id = parts[2]; category = str(data.get("category", "")).strip(); subject = str(data.get("subject", "")).strip(); state = str(data.get("state", "")).strip(); notes = str(data.get("notes", "")).strip()
                if category not in {"character", "prop", "timeline", "permission", "location"} or not subject: self.send_json({"error": "连续性类别或主体无效"}, 400); return
                try:
                    first = int(data["firstEpisode"]) if str(data.get("firstEpisode", "")).strip() else None; last = int(data["lastEpisode"]) if str(data.get("lastEpisode", "")).strip() else None
                except (TypeError, ValueError): self.send_json({"error": "集数必须是整数"}, 400); return
                if first is not None and last is not None and last < first: self.send_json({"error": "结束集不能早于起始集"}, 400); return
                entry_id = "c-" + uuid.uuid4().hex[:12]; stamp = now_iso()
                with db() as conn:
                    if not conn.execute("SELECT 1 FROM projects WHERE id=?", (project_id,)).fetchone(): self.send_json({"error": "项目不存在"}, 404); return
                    conn.execute("INSERT INTO continuity_entries(id,project_id,category,subject,state,first_episode,last_episode,source_version_id,status,notes,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)", (entry_id, project_id, category, subject, state, first, last, None, "active", notes, stamp, stamp))
                    conn.execute("INSERT INTO activity_logs(project_id,kind,message,created_at) VALUES(?,?,?,?)", (project_id, "continuity", f"新增连续性条目：{subject}", stamp))
                self.send_json({"entry": {"id": entry_id, "category": category, "subject": subject, "state": state, "first_episode": first, "last_episode": last, "source_version_id": None, "status": "active", "notes": notes, "created_at": stamp, "updated_at": stamp}}, 201); return
            if len(parts) == 5 and parts[:2] == ["api", "projects"] and parts[3:] == ["continuity", "extract"]:
                data = self.read_json(); project_id = parts[2]; version_id = str(data.get("versionId", ""))
                with db() as conn: version = conn.execute("SELECT * FROM draft_versions WHERE id=? AND project_id=?", (version_id, project_id)).fetchone()
                if not version: self.send_json({"error": "请选择当前项目的正文版本"}, 400); return
                analysis = analyze_script(version["content"]); char_eps, location_eps = {}, {}
                for episode in analysis["episodes"]:
                    for scene in episode["scenes"]:
                        for name in scene["characters"]: char_eps.setdefault(name, set()).add(episode["number"])
                        location_match = re.match(r"^\d+\s*[-－—]\s*\d+\s+\S+\s+\S+\s+(.+)$", scene["heading"]); location = location_match.group(1) if location_match else scene["heading"]
                        location_eps.setdefault(location, set()).add(episode["number"])
                stamp = now_iso(); entries = []
                with db() as conn:
                    conn.execute("DELETE FROM continuity_entries WHERE project_id=? AND source_version_id=? AND category IN ('character','location')", (project_id, version_id))
                    for category, mapping in (("character", char_eps), ("location", location_eps)):
                        for subject, eps in mapping.items():
                            entry_id = "c-" + uuid.uuid4().hex[:12]; first, last = min(eps), max(eps); state = f"正文识别：出现于{len(eps)}集"
                            conn.execute("INSERT INTO continuity_entries(id,project_id,category,subject,state,first_episode,last_episode,source_version_id,status,notes,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)", (entry_id, project_id, category, subject, state, first, last, version_id, "auto", "由场景标题与人物行自动提取，需人工补充状态变化", stamp, stamp)); entries.append(entry_id)
                    conn.execute("INSERT INTO activity_logs(project_id,kind,message,created_at) VALUES(?,?,?,?)", (project_id, "continuity", f"从{version['label']}提取连续性账本：{len(entries)}条", stamp))
                self.send_json({"ok": True, "count": len(entries), "versionId": version_id}, 201); return
            if len(parts) == 4 and parts[:2] == ["api", "projects"] and parts[3] == "sources":
                data = self.read_json(); name = str(data.get("name", "")).strip(); mime = str(data.get("mime", "application/octet-stream"))[:120]; encoded = str(data.get("dataBase64", ""))
                if not name or len(name) > 180 or not encoded: self.send_json({"error": "文件名和文件内容不能为空"}, 400); return
                try: raw = base64.b64decode(encoded, validate=True)
                except Exception: self.send_json({"error": "文件编码无效"}, 400); return
                if len(raw) > 10 * 1024 * 1024: self.send_json({"error": "单个素材不能超过10MB"}, 413); return
                project_id = parts[2]; source_id = "s-" + uuid.uuid4().hex[:12]; stamp = now_iso(); safe_name = "".join(c if c.isalnum() or c in "._-" else "_" for c in name)
                folder = STORAGE_ROOT / project_id / "sources" / source_id; folder.mkdir(parents=True, exist_ok=True); path = folder / safe_name; path.write_bytes(raw)
                digest = hashlib.sha256(raw).hexdigest(); extracted, status = extract_source_text(path); object_key = str(path.relative_to(ROOT)).replace("\\", "/")
                with db() as conn:
                    if not conn.execute("SELECT 1 FROM projects WHERE id=?", (project_id,)).fetchone(): self.send_json({"error": "项目不存在"}, 404); return
                    conn.execute("INSERT INTO source_documents(id,project_id,name,mime,object_key,byte_size,sha256,extracted_text,extraction_status,created_at) VALUES(?,?,?,?,?,?,?,?,?,?)", (source_id, project_id, name, mime, object_key, len(raw), digest, extracted, status, stamp))
                    conn.execute("INSERT INTO activity_logs(project_id,kind,message,created_at) VALUES(?,?,?,?)", (project_id, "source", f"导入素材：{name}（{status}）", stamp)); conn.execute("UPDATE projects SET updated_at=? WHERE id=?", (stamp, project_id))
                self.send_json({"source": {"id": source_id, "name": name, "mime": mime, "byte_size": len(raw), "sha256": digest, "extraction_status": status, "created_at": stamp}}, 201); return
        except (ValueError, json.JSONDecodeError) as exc: self.send_json({"error": f"请求格式错误：{exc}"}, 400); return
        self.send_json({"error": "接口不存在"}, 404)

    def do_PUT(self):
        path = urlparse(self.path).path.rstrip("/"); parts = self.parts(path)
        try:
            if len(parts) == 5 and parts[:2] == ["api", "projects"] and parts[3] == "assets" and parts[4] in ASSET_TYPES:
                data = self.read_json(); content = str(data.get("content", "")); stamp = now_iso()
                with db() as conn:
                    if not conn.execute("SELECT 1 FROM projects WHERE id=?", (parts[2],)).fetchone(): self.send_json({"error": "项目不存在"}, 404); return
                    conn.execute("""INSERT INTO story_assets(project_id,asset_type,content,version,updated_at) VALUES(?,?,?,1,?) ON CONFLICT(project_id,asset_type) DO UPDATE SET content=excluded.content,version=story_assets.version+1,updated_at=excluded.updated_at""", (parts[2], parts[4], content, stamp))
                    conn.execute("INSERT INTO activity_logs(project_id,kind,message,created_at) VALUES(?,?,?,?)", (parts[2], "asset", f"更新项目资料：{parts[4]}", stamp)); conn.execute("UPDATE projects SET updated_at=? WHERE id=?", (stamp, parts[2])); row = conn.execute("SELECT content,version,updated_at FROM story_assets WHERE project_id=? AND asset_type=?", (parts[2], parts[4])).fetchone()
                self.send_json({"asset": dict(row)}); return
        except (ValueError, json.JSONDecodeError) as exc: self.send_json({"error": f"请求格式错误：{exc}"}, 400); return
        self.send_json({"error": "接口不存在"}, 404)

    def do_DELETE(self):
        path = urlparse(self.path).path.rstrip("/"); parts = self.parts(path)
        if len(parts) == 3 and parts[:2] == ["api", "continuity"]:
            with db() as conn:
                row = conn.execute("SELECT project_id,subject FROM continuity_entries WHERE id=?", (parts[2],)).fetchone()
                if not row: self.send_json({"error": "连续性条目不存在"}, 404); return
                conn.execute("DELETE FROM continuity_entries WHERE id=?", (parts[2],)); conn.execute("INSERT INTO activity_logs(project_id,kind,message,created_at) VALUES(?,?,?,?)", (row["project_id"], "continuity", f"删除连续性条目：{row['subject']}", now_iso()))
            self.send_json({"ok": True}); return
        self.send_json({"error": "接口不存在"}, 404)


if __name__ == "__main__":
    init_db(); server = ThreadingHTTPServer(("127.0.0.1", 8786), Handler)
    print("剧本创作总控台：http://127.0.0.1:8786"); print(f"本地数据：{DB_PATH}"); server.serve_forever()
